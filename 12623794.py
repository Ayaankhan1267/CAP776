from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment
from datetime import datetime
import math
import os

class InvalidActivityDataError(Exception):
    """Raised when daily activity data is invalid."""
    pass

class DailyActivity:

    MINUTES_IN_DAY = 1440

    def __init__(
        self,
        date,
        sleep,
        fitness,
        study,
        coding,
        class_time,
        classes_attended,
        other_activities,
        feeling,
        satisfaction,
        energy,
        notes=""
    ):

        self.date = date
        self.sleep = sleep
        self.fitness = fitness
        self.study = study
        self.coding = coding
        self.class_time = class_time
        self.classes_attended = classes_attended
        self.other_activities = other_activities
        self.feeling = feeling
        self.satisfaction = satisfaction
        self.energy = energy
        self.notes = notes

        self.validate()

    def validate(self):

        time_values = [
            self.sleep,
            self.fitness,
            self.study,
            self.coding,
            self.class_time,
            self.other_activities
        ]

        for value in time_values:

            if value < 0:
                raise InvalidActivityDataError(
                    "Activity time cannot be negative."
                )

        if self.classes_attended < 0:
            raise InvalidActivityDataError(
                "Classes attended cannot be negative."
            )

        if self.total_tracked() > self.MINUTES_IN_DAY:
            raise InvalidActivityDataError(
                "Total tracked time cannot be greater than 1440 minutes."
            )

        if self.feeling_score() is None:
            raise InvalidActivityDataError(
                "Invalid feeling value."
            )

        if self.satisfaction_score() is None:
            raise InvalidActivityDataError(
                "Invalid satisfaction value."
            )

        if self.energy_score() is None:
            raise InvalidActivityDataError(
                "Invalid energy value."
            )

    def total_tracked(self):

        return (
            self.sleep
            + self.fitness
            + self.study
            + self.coding
            + self.class_time
            + self.other_activities
        )

    def free_time(self):

        return self.MINUTES_IN_DAY - self.total_tracked()

    def feeling_score(self):

        scores = {
            "Excellent": 5,
            "Good": 4,
            "Neutral": 3,
            "Low": 2,
            "Stressed": 1
        }

        return scores.get(
            str(self.feeling).strip(),
            None
        )

    def satisfaction_score(self):

        scores = {
            "Very Satisfied": 5,
            "Satisfied": 4,
            "Neutral": 3,
            "Unsatisfied": 2,
            "Very Unsatisfied": 1
        }

        return scores.get(
            str(self.satisfaction).strip(),
            None
        )

    def energy_score(self):

        scores = {
            "High": 3,
            "Medium": 2,
            "Low": 1
        }

        return scores.get(
            str(self.energy).strip(),
            None
        )

    def experience_score(self):

        feeling = self.feeling_score()
        satisfaction = self.satisfaction_score()
        energy = self.energy_score()

        if (
            feeling is None
            or satisfaction is None
            or energy is None
        ):
            return None

        return (
            feeling
            + satisfaction
            + energy
        ) / 3

def read_excel_data(filename):

    workbook = load_workbook(
        filename,
        data_only=True
    )

    if "Daily Log" not in workbook.sheetnames:
        raise KeyError(
            "Daily Log sheet was not found."
        )

    sheet = workbook["Daily Log"]

    header_row = 5

    headers = {}

    for cell in sheet[header_row]:

        if cell.value is not None:
            headers[str(cell.value).strip()] = cell.column

    required_columns = [
        "Date",
        "Sleep (min)",
        "Fitness (min)",
        "Study (min)",
        "Coding (min)",
        "Class (min)",
        "Classes Attended",
        "Other Activities (min)",
        "Day's Feeling",
        "Satisfaction Level",
        "Energy Level",
        "Notes"
    ]

    for column in required_columns:

        if column not in headers:
            raise KeyError(
                f"Required column missing: {column}"
            )

    activities = []
    invalid_records = 0

    for row_number in range(
        header_row + 1,
        sheet.max_row + 1
    ):

        date_value = sheet.cell(
            row=row_number,
            column=headers["Date"]
        ).value

        if date_value is None:
            continue

        try:

            if isinstance(date_value, datetime):
                date = date_value.date()

            else:
                date = datetime.strptime(
                    str(date_value),
                    "%Y-%m-%d"
                ).date()

            sleep = float(
                sheet.cell(
                    row=row_number,
                    column=headers["Sleep (min)"]
                ).value
            )

            fitness = float(
                sheet.cell(
                    row=row_number,
                    column=headers["Fitness (min)"]
                ).value
            )

            study = float(
                sheet.cell(
                    row=row_number,
                    column=headers["Study (min)"]
                ).value
            )

            coding = float(
                sheet.cell(
                    row=row_number,
                    column=headers["Coding (min)"]
                ).value
            )

            class_time = float(
                sheet.cell(
                    row=row_number,
                    column=headers["Class (min)"]
                ).value
            )

            classes_attended = float(
                sheet.cell(
                    row=row_number,
                    column=headers["Classes Attended"]
                ).value
            )

            other_activities = float(
                sheet.cell(
                    row=row_number,
                    column=headers["Other Activities (min)"]
                ).value
            )

            feeling = sheet.cell(
                row=row_number,
                column=headers["Day's Feeling"]
            ).value

            satisfaction = sheet.cell(
                row=row_number,
                column=headers["Satisfaction Level"]
            ).value

            energy = sheet.cell(
                row=row_number,
                column=headers["Energy Level"]
            ).value

            notes = sheet.cell(
                row=row_number,
                column=headers["Notes"]
            ).value

            if notes is None:
                notes = ""

            else:
                notes = str(notes)

            activity = DailyActivity(
                date,
                sleep,
                fitness,
                study,
                coding,
                class_time,
                classes_attended,
                other_activities,
                feeling,
                satisfaction,
                energy,
                notes
            )

            activities.append(activity)

        except (
            ValueError,
            TypeError,
            InvalidActivityDataError
        ):

            invalid_records += 1

    workbook.close()

    return activities, invalid_records

def average(values):

    if len(values) == 0:
        return 0

    return sum(values) / len(values)

def calculate_summary(activities):

    return {

        "Average Sleep":
            average([
                activity.sleep
                for activity in activities
            ]),

        "Average Fitness":
            average([
                activity.fitness
                for activity in activities
            ]),

        "Average Study":
            average([
                activity.study
                for activity in activities
            ]),

        "Average Coding":
            average([
                activity.coding
                for activity in activities
            ]),

        "Average Class":
            average([
                activity.class_time
                for activity in activities
            ]),

        "Average Other Activities":
            average([
                activity.other_activities
                for activity in activities
            ]),

        "Average Free / Unaccounted Time":
            average([
                activity.free_time()
                for activity in activities
            ]),

        "Valid Days":
            len(activities)
    }

def calculate_indices(activities):

    valid_days = len(activities)

    if valid_days == 0:
        raise InvalidActivityDataError(
            "No valid activity records available."
        )

    TPI = average([
        activity.coding
        for activity in activities
    ])

    AAI = average([
        activity.study + activity.class_time
        for activity in activities
    ])

    PhAI = average([
        activity.fitness
        for activity in activities
    ])

    SRI = average([
        activity.sleep
        for activity in activities
    ])

    ABI = average([
        activity.free_time()
        for activity in activities
    ])

    TUI = average([
        activity.total_tracked()
        for activity in activities
    ])

    EI = average([
        activity.experience_score()
        for activity in activities
    ])

    expected_days = 40

    DCI = (
        valid_days / expected_days
    ) * 100

    PAI = (
        0.15 * TPI
        + 0.20 * AAI
        + 0.15 * PhAI
        + 0.20 * SRI
        + 0.15 * TUI
        + 0.10 * EI
        + 0.05 * DCI
    )

    return {

        "Expected Days": expected_days,

        "Valid Days": valid_days,

        "TPI": TPI,

        "AAI": AAI,

        "PhAI": PhAI,

        "SRI": SRI,

        "ABI": ABI,

        "TUI": TUI,

        "EI": EI,

        "DCI": DCI,

        "PAI": PAI
    }

def pearson_correlation(x, y):

    if len(x) != len(y):
        return None

    if len(x) < 2:
        return None

    mean_x = average(x)
    mean_y = average(y)

    numerator = 0
    sum_x = 0
    sum_y = 0

    for i in range(len(x)):

        difference_x = x[i] - mean_x
        difference_y = y[i] - mean_y

        numerator += (
            difference_x * difference_y
        )

        sum_x += difference_x ** 2
        sum_y += difference_y ** 2

    denominator = math.sqrt(
        sum_x * sum_y
    )

    if denominator == 0:
        return 0

    return numerator / denominator

def calculate_correlations(activities):

    sleep = [
        activity.sleep
        for activity in activities
    ]

    energy = [
        activity.energy_score()
        for activity in activities
    ]

    study = [
        activity.study
        for activity in activities
    ]

    satisfaction = [
        activity.satisfaction_score()
        for activity in activities
    ]

    coding = [
        activity.coding
        for activity in activities
    ]

    return {

        "Sleep ↔ Energy":
            pearson_correlation(
                sleep,
                energy
            ),

        "Study ↔ Satisfaction":
            pearson_correlation(
                study,
                satisfaction
            ),

        "Coding ↔ Energy":
            pearson_correlation(
                coding,
                energy
            )
    }

def correlation_note(value):

    if value is None:
        return "Not enough data"

    if value >= 0.7:
        return "Strong positive relationship"

    elif value >= 0.3:
        return "Moderate positive relationship"

    elif value > -0.3:
        return "Weak or little relationship"

    elif value > -0.7:
        return "Moderate negative relationship"

    else:
        return "Strong negative relationship"

def print_report(
    summary,
    indices,
    correlations,
    invalid_records
):

    print()
    print("=" * 60)
    print("CAP776 - PERSONAL ACTIVITY INTELLIGENCE REPORT")
    print("=" * 60)

    print("\nACTIVITY SUMMARY")
    print("-" * 60)

    print(
        f"Expected Days: "
        f"{indices['Expected Days']}"
    )

    print(
        f"Valid Days: "
        f"{indices['Valid Days']}"
    )

    print(
        f"Invalid / Excluded Records: "
        f"{invalid_records}"
    )

    print(
        f"Average Sleep: "
        f"{summary['Average Sleep']:.2f} min/day"
    )

    print(
        f"Average Fitness: "
        f"{summary['Average Fitness']:.2f} min/day"
    )

    print(
        f"Average Study: "
        f"{summary['Average Study']:.2f} min/day"
    )

    print(
        f"Average Coding: "
        f"{summary['Average Coding']:.2f} min/day"
    )

    print(
        f"Average Class: "
        f"{summary['Average Class']:.2f} min/day"
    )

    print(
        f"Average Other Activities: "
        f"{summary['Average Other Activities']:.2f} min/day"
    )

    print(
        f"Average Free/Unaccounted: "
        f"{summary['Average Free / Unaccounted Time']:.2f} min/day"
    )

    print("\nINDEX VALUES")
    print("-" * 60)

    for name in [
        "TPI",
        "AAI",
        "PhAI",
        "SRI",
        "ABI",
        "TUI",
        "EI",
        "DCI",
        "PAI"
    ]:

        if name == "DCI":
            print(
                f"{name}: "
                f"{indices[name]:.2f}%"
            )

        else:
            print(
                f"{name}: "
                f"{indices[name]:.2f}"
            )

    print("\nCORRELATION OBSERVATIONS")
    print("-" * 60)

    for relationship, value in correlations.items():

        if value is None:
            print(
                f"{relationship}: "
                f"Not available"
            )

        else:
            print(
                f"{relationship}: "
                f"{value:.2f} - "
                f"{correlation_note(value)}"
            )

def save_results(
    activities,
    summary,
    indices,
    correlations,
    invalid_records,
    filename
):

    workbook = Workbook()

    daily_sheet = workbook.active
    daily_sheet.title = "Calculated Daily Data"

    headers = [
        "Date",
        "Sleep (min)",
        "Fitness (min)",
        "Study (min)",
        "Coding (min)",
        "Class (min)",
        "Classes Attended",
        "Other Activities (min)",
        "Total Tracked (min)",
        "Free/Unaccounted (min)",
        "Day's Feeling",
        "Feeling Score",
        "Satisfaction Level",
        "Satisfaction Score",
        "Energy Level",
        "Energy Score",
        "Experience Score",
        "Notes"
    ]

    daily_sheet.append(headers)

    for cell in daily_sheet[1]:
        cell.font = Font(bold=True)

    for activity in activities:

        daily_sheet.append([

            activity.date,

            activity.sleep,

            activity.fitness,

            activity.study,

            activity.coding,

            activity.class_time,

            activity.classes_attended,

            activity.other_activities,

            activity.total_tracked(),

            activity.free_time(),

            activity.feeling,

            activity.feeling_score(),

            activity.satisfaction,

            activity.satisfaction_score(),

            activity.energy,

            activity.energy_score(),

            activity.experience_score(),

            activity.notes
        ])

    result_sheet = workbook.create_sheet(
        "Project Results"
    )

    result_sheet.append([
        "Parameter",
        "Value"
    ])

    result_sheet["A1"].font = Font(bold=True)
    result_sheet["B1"].font = Font(bold=True)

    results = [

        (
            "Expected Days",
            indices["Expected Days"]
        ),

        (
            "Valid Days",
            indices["Valid Days"]
        ),

        (
            "Invalid / Excluded Records",
            invalid_records
        ),

        (
            "Average Sleep (min/day)",
            summary["Average Sleep"]
        ),

        (
            "Average Fitness (min/day)",
            summary["Average Fitness"]
        ),

        (
            "Average Study (min/day)",
            summary["Average Study"]
        ),

        (
            "Average Coding (min/day)",
            summary["Average Coding"]
        ),

        (
            "Average Class (min/day)",
            summary["Average Class"]
        ),

        (
            "Average Other Activities (min/day)",
            summary["Average Other Activities"]
        ),

        (
            "Average Free / Unaccounted Time (min/day)",
            summary["Average Free / Unaccounted Time"]
        ),

        (
            "TPI",
            indices["TPI"]
        ),

        (
            "AAI",
            indices["AAI"]
        ),

        (
            "PhAI",
            indices["PhAI"]
        ),

        (
            "SRI",
            indices["SRI"]
        ),

        (
            "ABI",
            indices["ABI"]
        ),

        (
            "TUI",
            indices["TUI"]
        ),

        (
            "EI",
            indices["EI"]
        ),

        (
            "DCI",
            indices["DCI"]
        ),

        (
            "PAI",
            indices["PAI"]
        ),

        (
            "Sleep ↔ Energy",
            correlations["Sleep ↔ Energy"]
        ),

        (
            "Study ↔ Satisfaction",
            correlations["Study ↔ Satisfaction"]
        ),

        (
            "Coding ↔ Energy",
            correlations["Coding ↔ Energy"]
        )
    ]

    for parameter, value in results:

        result_sheet.append([
            parameter,
            value
        ])

    observation_sheet = workbook.create_sheet(
        "Observations"
    )

    observation_sheet.append([
        "Relationship",
        "Correlation",
        "Observation"
    ])

    for cell in observation_sheet[1]:
        cell.font = Font(bold=True)

    for relationship, value in correlations.items():

        observation_sheet.append([

            relationship,

            value,

            correlation_note(value)
        ])

    for sheet in workbook.worksheets:

        for column in sheet.columns:

            max_length = 0

            for cell in column:

                if cell.value is not None:

                    length = len(
                        str(cell.value)
                    )

                    if length > max_length:
                        max_length = length

            sheet.column_dimensions[
                column[0].column_letter
            ].width = min(
                max_length + 2,
                40
            )

        for row in sheet.iter_rows():

            for cell in row:
                cell.alignment = Alignment(
                    vertical="center"
                )

    workbook.save(filename)

def main():

    input_file = "12623794.xlsx"

    output_file = "12623794_results.xlsx"

    try:

        print("Reading Excel file...")

        activities, invalid_records = (
            read_excel_data(input_file)
        )

        if len(activities) == 0:

            print(
                "No valid activity records found."
            )

            return

        summary = calculate_summary(
            activities
        )

        indices = calculate_indices(
            activities
        )

        correlations = calculate_correlations(
            activities
        )

        print_report(
            summary,
            indices,
            correlations,
            invalid_records
        )

        save_results(
            activities,
            summary,
            indices,
            correlations,
            invalid_records,
            output_file
        )

        print()
        print("=" * 60)
        print(
            f"Results saved to: {output_file}"
        )
        print("=" * 60)

    except FileNotFoundError:

        print(
            f"Error: {input_file} was not found."
        )

        print(
            "Make sure the Excel file is in "
            "the same folder as this Python file."
        )

    except PermissionError:

        print(
            "Permission error while accessing "
            "the Excel file."
        )

        print(
            "Close the Excel workbook if it is open "
            "and run the program again."
        )

    except KeyError as error:

        print(
            f"Excel structure error: {error}"
        )

    except InvalidActivityDataError as error:

        print(
            f"Invalid activity data: {error}"
        )

    except Exception as error:

        print(
            f"Unexpected error: {error}"
        )

if __name__ == "__main__":
    main()